"""Add ActiveListings / LastListingCreated / HasConnection columns to the
professional-customers Excel by joining DB metrics on ImmoConnect_OrganizationId
(with fallback on OrganizationNumber / EstablishmentNumber)."""
import json
import openpyxl

XLSX = r"products\spotto\drafts\customers-clean-up\professional-customers-2026-06-19.xlsx"
METRICS = r"C:\Users\bartp\.claude\projects\c--Users-bartp-source-repos-spotto-os\8afa04c0-eb0d-4386-a149-e047b7766e57\tool-results\mcp-mssql-read_data-1782132186863.txt"


def norm(v):
    return str(v).strip().lower() if v not in (None, "") else None


# --- load DB metrics, build lookup tables keyed by each available identifier ---
with open(METRICS, encoding="utf-8") as f:
    rows = json.load(f)["data"]

by_immo, by_org, by_est = {}, {}, {}
for r in rows:
    rec = (r["ActiveListings"], r["LastListingCreated"], r["HasConnection"])
    if norm(r["ImmoId"]):
        by_immo[norm(r["ImmoId"])] = rec
    if norm(r["OrgNum"]):
        by_org.setdefault(norm(r["OrgNum"]), rec)
    if norm(r["EstNum"]):
        by_est.setdefault(norm(r["EstNum"]), rec)

print(f"DB metric rows: {len(rows)}")

# --- enrich the workbook ---
wb = openpyxl.load_workbook(XLSX)
ws = wb["Professional customers"]
header = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(header)}
base = ws.max_column

ws.cell(1, base + 1, "ActiveListings")
ws.cell(1, base + 2, "LastListingCreated")
ws.cell(1, base + 3, "HasConnection")

matched = unmatched = 0
for row in range(2, ws.max_row + 1):
    immo = norm(ws.cell(row, col["ImmoConnect_OrganizationId"]).value)
    org = norm(ws.cell(row, col["OrganizationNumber"]).value)
    est = norm(ws.cell(row, col["EstablishmentNumber"]).value)
    rec = (immo and by_immo.get(immo)) or (org and by_org.get(org)) or (est and by_est.get(est))
    if rec:
        matched += 1
        ws.cell(row, base + 1, rec[0])
        ws.cell(row, base + 2, rec[1])
        ws.cell(row, base + 3, "Yes" if rec[2] else "No")
    else:
        unmatched += 1
        ws.cell(row, base + 1, None)
        ws.cell(row, base + 2, None)
        ws.cell(row, base + 3, None)

wb.save(XLSX)
print(f"rows enriched: matched={matched}, unmatched={unmatched}")
