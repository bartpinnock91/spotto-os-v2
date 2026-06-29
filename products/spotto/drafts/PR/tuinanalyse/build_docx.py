# -*- coding: utf-8 -*-
"""Build a Word version of the tuinanalyse press prep document from the markdown."""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = sys.argv[1] if len(sys.argv) > 1 else "persbericht.md"
OUT = SRC.rsplit(".", 1)[0] + ".docx"

doc = Document()

# --- base styling ---
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

ACCENT = RGBColor(0x1F, 0x6F, 0x43)  # spotto-ish green for headings


def add_runs(paragraph, text):
    """Render **bold** segments within a markdown line."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part == "":
            continue
        run = paragraph.add_run(part)
        if i % 2 == 1:
            run.bold = True


def add_table(lines):
    """lines: list of markdown table rows (header, separator, data...)."""
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    header, data = rows[0], rows[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for j, cell in enumerate(header):
        p = table.rows[0].cells[j].paragraphs[0]
        add_runs(p, cell)
        for r in p.runs:
            r.bold = True
    for drow in data:
        cells = table.add_row().cells
        for j, val in enumerate(drow):
            add_runs(cells[j].paragraphs[0], val)
    doc.add_paragraph()


with open(SRC, encoding="utf-8") as fh:
    raw = fh.read().splitlines()

i = 0
while i < len(raw):
    line = raw[i].rstrip()

    if not line.strip():
        i += 1
        continue
    if line.strip() == "---":
        i += 1
        continue

    # tables: a block of lines starting with "|"
    if line.lstrip().startswith("|"):
        block = []
        while i < len(raw) and raw[i].lstrip().startswith("|"):
            block.append(raw[i])
            i += 1
        add_table(block)
        continue

    # headings
    m = re.match(r"^(#{1,6})\s+(.*)$", line)
    if m:
        level = len(m.group(1))
        text = m.group(2)
        if level == 1:
            h = doc.add_heading(level=0)
            add_runs(h, text)
        else:
            h = doc.add_heading(level=min(level, 4))
            h.runs and None
            for r in h.runs:
                r.font.color.rgb = ACCENT
            if not h.runs:
                add_runs(h, text)
            else:
                h.runs[0].text = ""
                add_runs(h, text)
                for r in h.runs:
                    r.font.color.rgb = ACCENT
        i += 1
        continue

    # bullet list
    if line.lstrip().startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, line.lstrip()[2:])
        i += 1
        continue

    # takeaway arrow lines -> indented
    if line.lstrip().startswith("→"):
        p = doc.add_paragraph()
        add_runs(p, line.strip())
        p.paragraph_format.left_indent = Pt(12)
        i += 1
        continue

    # italic notes (_..._)
    mi = re.match(r"^_(.*)_$", line.strip())
    if mi:
        p = doc.add_paragraph()
        run = p.add_run(mi.group(1))
        run.italic = True
        run.font.size = Pt(9)
        i += 1
        continue

    # plain paragraph
    p = doc.add_paragraph()
    add_runs(p, line)
    i += 1

doc.save(OUT)
print("Saved", OUT)
