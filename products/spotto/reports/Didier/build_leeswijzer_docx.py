# Genereert leeswijzer-weekrapport.docx (Word-versie van leeswijzer-weekrapport.md).
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT
    return h


def para(doc, text, italic=False, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for c, h in enumerate(headers):
        cell = t.rows[0].cells[c]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row):
            t.rows[r].cells[c].text = v
    if widths:
        for c, w in enumerate(widths):
            for row in t.rows:
                row.cells[c].width = Cm(w)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph()
    return t


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

heading(doc, "Leeswijzer — Spotto weekrapport 2026", 0)
para(doc, "Periode: week 1 (29 dec 2025 – 4 jan 2026) t/m week 24 (onvolledig, t/m 12 juni).")
para(doc, "Weken = ISO-weken, maandag t/m zondag. Pandtypes: huis / appartement / overig "
          "(overig = grond, garage, commercieel, kamers, onbekend). Koop/huur overal gesplitst.")

heading(doc, "Sectie 1 — Aanbod (publicaties)", 1)
para(doc, "Telt publicaties (advertenties): een pand dat opnieuw gepubliceerd wordt of van "
          "kantoor wisselt, telt opnieuw mee.")
table(doc, ["Tabblad", "Definitie"], [
    ["Nieuw gepubliceerd", "Publicaties die die week voor het eerst online kwamen"],
    ["Online aanbod", "Stand op zondagavond + aantal kantoren met ≥1 pand online"],
    ["Offline gehaald", "Publicaties die die week offline gingen (verkocht/verhuurd/"
     "teruggetrokken — reden is niet gekend)"],
    ["Vraagprijzen huis / appartement", "Mediaan en gemiddelde vraagprijs van nieuwe "
     "publicaties met prijs; N = aantal met prijs. Gebruik de mediaan"],
    ["EPC huis / appartement", "Labelverdeling nieuwe publicaties; gewesten samengevoegd per letter"],
    ["Nieuwbouw vs bestaand", "Projectpublicatie = 1 publicatie, kan meerdere units omvatten"],
], widths=[5.5, 11.5])

heading(doc, "Sectie 2 — Objecten (unieke panden)", 1)
para(doc, "Telt unieke panden (adresniveau). Verschil met sectie 1 = herpublicaties/kantoorwissels.")
table(doc, ["Tabblad", "Definitie"], [
    ["Op de markt gekomen", "Unieke panden die een marktperiode startten; eerste keer vs herpublicatie"],
    ["Doorlooptijd", "Mediaan aantal dagen op de markt van panden die die week van de markt "
     "gingen; gemengde koop/huur-periodes tellen als koop"],
    ["Prijsverlagingen", "Aantal en mediane daling (€ en %); enkel huis + appartement"],
    ["Startprijs huis / appartement", "Mediaan start- vs huidige prijs per weekcohorte + % ooit verlaagd"],
], widths=[5.5, 11.5])
para(doc, "Nieuwbouwunits zonder busnummer tellen als één pand. Cijfers lopen ±2-3 dagen "
          "achter; week 24 is hier extra onvolledig.")

heading(doc, "Sectie 3 — Bezoekers & zoekers", 1)
table(doc, ["Tabblad", "Definitie"], [
    ["Bezoekers", "Unieke bezoekers, nieuwe bezoekers, sessies, engagement rate"],
    ["Zoekers koop vs huur", "Unieke bezoekers die die week ≥1 koop- resp. huurpandpagina "
     "bekeken. Overlap mogelijk: niet optellen. Eenzelfde persoon telt elke week opnieuw"],
    ["Pandweergaves", "Overzichts- en detailpagina's"],
    ["Zoekopdrachten", "Uitgevoerde zoekopdrachten + unieke zoekende gebruikers"],
    ["Verkeersbronnen", "Sessies per kanaal"],
], widths=[5.5, 11.5])
para(doc, "Vaste notitie: deze cijfers onderschatten de werkelijkheid met ±33% "
          "(cookieweigeraars en adblockers). Absolute aantallen zijn ondergrenzen; trends en "
          "verhoudingen zijn betrouwbaar. Bezoeker = browser/toestel, geen persoon.", italic=True)

heading(doc, "Sectie 4 — Leads & interactie", 1)
table(doc, ["Tabblad", "Definitie"], [
    ["Contactaanvragen", "Verzonden contactformulieren op panden, gesplitst infovraag / bezoekaanvraag"],
    ["Klikgedrag contact", "Kliks op bel-makelaar en formulier-CTA's"],
    ["Registraties", "Nieuwe accounts; waarvan makelaar; met nieuwsbriefinschrijving"],
    ["Favorieten", "Bewaarde panden"],
    ["Partnerbanners", "Weergaves, bereikte gebruikers en kliks op partneradvertenties"],
], widths=[5.5, 11.5])

heading(doc, "Let op bij het lezen (gekende gebeurtenissen, géén markttendensen)", 1)
table(doc, ["Periode", "Wat", "Effect"], [
    ["Week 6–9 (feb)", "Koppeling Dewaele stuurde te veel publicaties door",
     "Piek in nieuw én offline (sectie 1); online aanbod en objectcijfers (sectie 2) bleven correct"],
    ["Vanaf week 9 (maart)", "Betaalde zoekadvertenties bewust teruggeschroefd",
     "Daling bezoekers en registraties is deels hierdoor, geen marktafkoeling"],
    ["Week 12 (18 maart)", "Contactformulier aangepast: dropdown volgt voortaan de geklikte knop",
     "Verhouding infovraag/bezoekaanvraag pas vanaf week 13 betrouwbaar; totalen wel doorlopend bruikbaar"],
    ["Week 14 (1 april)", "Nieuwe mobiele actiebalk (favoriet | bezoek | bel | pandscore)",
     "Bel-makelaar-kliks dalen blijvend; verschuiving naar bezoekaanvragen, geen leadverlies"],
    ["Week 23 (2 juni)", "Commerciële portefeuille Ceusters toegevoegd (±900 huurpanden)",
     "Huuraanbod 'overig' structureel +0,9k; eenmalige piek nieuwe huurpublicaties. Residentieel onaangetast"],
    ["Week 24", "Lopende week",
     "Alle cijfers onvolledig (t/m vr 12 juni); verkeersbronnen pas definitief na weekafsluiting"],
], widths=[3.5, 6.0, 7.5])

para(doc, "Algemeen: publicatiecijfers (sectie 1) en pandcijfers (sectie 2) beantwoorden "
          "verschillende vragen — hoeveel aanbod stond online vs hoeveel unieke panden kwamen "
          "op de markt. Het verschil is herpublicatie, geen dubbeltelling-fout.", bold=True)

out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, "leeswijzer-weekrapport.docx")
doc.save(out_path)
print(f"OK: {out_path}")
