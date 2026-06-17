# Bouwt het wekelijkse Spotto-rapport, sectie 2 (Objectniveau - Vergelijkingspanden).
# Bron: realestatecomparison-productiedatabank, schema databricks (ReferenceProperties*),
# opgehaald 2026-06-12. Objecten = unieke adressen (address_key).
# Weken = ISO-weken. Let op: Databricks-sync loopt 2-3 dagen achter; week 24 is dubbel onvolledig.
import datetime
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# (weekstart, objecten op de markt gekomen, waarvan eerste keer, mediaan dagen sinds vorige episode,
#  ter vergelijking: nieuwe publicaties die week uit sectie 1, koop+huur)
MARKT = [
    ("2025-12-29", 638, 474, 484, 1647), ("2026-01-05", 1477, 1190, 476, 4012),
    ("2026-01-12", 1366, 1109, 476, 3519), ("2026-01-19", 1583, 1296, 546, 4191),
    ("2026-01-26", 1508, 1260, 519.5, 5708), ("2026-02-02", 1437, 1135, 519.5, 5816),
    ("2026-02-09", 1601, 1299, 513, 8989), ("2026-02-16", 1646, 1283, 559, 9405),
    ("2026-02-23", 1586, 1323, 498, 11252), ("2026-03-02", 1738, 1411, 503, 5385),
    ("2026-03-09", 1557, 1294, 515, 3427), ("2026-03-16", 1705, 1397, 487.5, 3762),
    ("2026-03-23", 1826, 1530, 528.5, 3874), ("2026-03-30", 1716, 1393, 497, 4006),
    ("2026-04-06", 1315, 1071, 538.5, 2685), ("2026-04-13", 1550, 1261, 528, 3316),
    ("2026-04-20", 1702, 1406, 544.5, 4193), ("2026-04-27", 1490, 1222, 495.5, 3122),
    ("2026-05-04", 1890, 1575, 496, 4135), ("2026-05-11", 1336, 1090, 548, 2985),
    ("2026-05-18", 1713, 1407, 582.5, 4303), ("2026-05-25", 1378, 1125, 541, 3193),
    ("2026-06-01", 1775, 1388, 544, 4837), ("2026-06-08", 639, 509, 550, 3026),
]

# (weekstart, koop N, koop mediaan dagen, huur N, huur mediaan dagen) - van de markt gegaan
DOORLOOPTIJD = [
    ("2025-12-29", 689, 115, 195, 37), ("2026-01-05", 1103, 110, 395, 46),
    ("2026-01-12", 1598, 126, 508, 42.5), ("2026-01-19", 1142, 112.5, 438, 42.5),
    ("2026-01-26", 1065, 94, 455, 46), ("2026-02-02", 1088, 104, 438, 35.5),
    ("2026-02-09", 1189, 95, 455, 31), ("2026-02-16", 1269, 89, 460, 33),
    ("2026-02-23", 1292, 122.5, 486, 38), ("2026-03-02", 1233, 95, 399, 40),
    ("2026-03-09", 1157, 105, 447, 31), ("2026-03-16", 1253, 113, 545, 40),
    ("2026-03-23", 1257, 125, 466, 35), ("2026-03-30", 1435, 120, 503, 37),
    ("2026-04-06", 813, 105, 366, 40), ("2026-04-13", 1078, 87, 416, 34),
    ("2026-04-20", 1309, 101, 494, 37), ("2026-04-27", 1042, 80.5, 458, 35),
    ("2026-05-04", 1210, 86.5, 412, 37.5), ("2026-05-11", 957, 73, 381, 42),
    ("2026-05-18", 1459, 96, 533, 41), ("2026-05-25", 1044, 81, 435, 36),
    ("2026-06-01", 1491, 88, 468, 35), ("2026-06-08", 614, 79.5, 220, 37.5),
]

# (weekstart, koop: N, mediaan daling EUR, mediaan daling %, huur: idem) - huis+app
PRIJSVERLAGINGEN = [
    ("2025-12-29", 44, 14000, 4.92, 5, 85, 6.14), ("2026-01-05", 181, 15000, 3.99, 22, 50, 5.01),
    ("2026-01-12", 129, 15000, 5.02, 18, 77.5, 6.47), ("2026-01-19", 137, 20000, 5.26, 41, 55, 5.56),
    ("2026-01-26", 143, 16000, 5.01, 32, 72.5, 6.25), ("2026-02-02", 111, 16000, 5.03, 20, 82.5, 6.59),
    ("2026-02-09", 145, 15000, 5.02, 24, 100, 6.92), ("2026-02-16", 126, 20000, 5.56, 20, 97.5, 8.13),
    ("2026-02-23", 138, 16000, 4.39, 22, 100, 6.82), ("2026-03-02", 141, 15000, 4.78, 23, 75, 7.14),
    ("2026-03-09", 118, 19499, 4.86, 21, 100, 6.33), ("2026-03-16", 132, 16000, 5.03, 25, 50, 4.17),
    ("2026-03-23", 172, 20000, 5.10, 21, 100, 6.11), ("2026-03-30", 140, 15000, 4.38, 28, 70, 6.69),
    ("2026-04-06", 120, 15500, 4.53, 20, 55, 5.67), ("2026-04-13", 165, 18000, 4.70, 25, 50, 4.80),
    ("2026-04-20", 159, 15000, 4.65, 19, 50, 5.06), ("2026-04-27", 136, 16000, 4.94, 25, 50, 5.56),
    ("2026-05-04", 172, 15000, 5.28, 17, 90, 7.33), ("2026-05-11", 119, 15000, 5.04, 21, 80, 6.45),
    ("2026-05-18", 156, 20000, 5.29, 22, 100, 5.71), ("2026-05-25", 144, 16000, 4.51, 13, 50, 5.71),
    ("2026-06-01", 184, 15134, 4.34, 23, 95, 6.25), ("2026-06-08", 64, 15500, 5.27, 6, 135, 9.48),
]

# (weekstart, koop: N, med startprijs, med huidige prijs, aandeel verlaagd, huur: idem) - HUIZEN
START_HUIS = [
    ("2025-12-29", 440, 395000, 395000, 0.0568, 52, 1100, 1100, 0.0385),
    ("2026-01-05", 1456, 420000, 419000, 0.0433, 152, 1200, 1200, 0.0395),
    ("2026-01-12", 1140, 417207.5, 417207.5, 0.0316, 137, 1195, 1190, 0.0438),
    ("2026-01-19", 1512, 398000, 398000, 0.0377, 154, 1200, 1195, 0.0195),
    ("2026-01-26", 2017, 399500, 399000, 0.0322, 115, 1100, 1100, 0.0174),
    ("2026-02-02", 2123, 439000, 439000, 0.0188, 150, 1100, 1100, 0.0067),
    ("2026-02-09", 3421, 439000, 440000, 0.0123, 127, 1125, 1125, 0.0236),
    ("2026-02-16", 3743, 447000, 447000, 0.0184, 142, 1250, 1250, 0.0211),
    ("2026-02-23", 4516, 449000, 449000, 0.0109, 108, 1250, 1250, 0.0278),
    ("2026-03-02", 1900, 418655, 416305, 0.0279, 153, 1100, 1100, 0.0196),
    ("2026-03-09", 1261, 425000, 425000, 0.0484, 94, 1100, 1100, 0.0106),
    ("2026-03-16", 1288, 419000, 419000, 0.0396, 134, 1200, 1200, 0.0299),
    ("2026-03-23", 1514, 435000, 434923.5, 0.0357, 132, 1200, 1175, 0.0227),
    ("2026-03-30", 1377, 428000, 427630, 0.0356, 132, 1200, 1200, 0.0303),
    ("2026-04-06", 975, 424340, 423000, 0.0379, 111, 1150, 1150, 0.0090),
    ("2026-04-13", 1228, 399900, 399615, 0.0432, 116, 1175, 1175, 0.0259),
    ("2026-04-20", 1539, 430000, 429500, 0.0325, 145, 1250, 1250, 0.0207),
    ("2026-04-27", 1190, 419500, 419500, 0.0269, 127, 1200, 1200, 0.0157),
    ("2026-05-04", 1585, 420000, 420000, 0.0246, 143, 1200, 1200, 0.0070),
    ("2026-05-11", 1205, 435000, 435000, 0.0149, 104, 1200, 1200, 0.0),
    ("2026-05-18", 1554, 429724, 429724, 0.0109, 135, 1200, 1200, 0.0),
    ("2026-05-25", 1266, 430000, 429250, 0.0087, 122, 1197.5, 1197.5, 0.0164),
    ("2026-06-01", 1454, 425000, 425000, 0.0062, 143, 1250, 1250, 0.0),
    ("2026-06-08", 533, 411420, 411420, 0.0, 50, 1100, 1100, 0.0),
]

# idem - APPARTEMENTEN
START_APP = [
    ("2025-12-29", 361, 299000, 299000, 0.0554, 147, 950, 950, 0.0136),
    ("2026-01-05", 1168, 299000, 299000, 0.0351, 473, 900, 900, 0.0402),
    ("2026-01-12", 1071, 301200, 300410, 0.0364, 440, 905, 900, 0.0636),
    ("2026-01-19", 1403, 299000, 299000, 0.0335, 437, 940, 935, 0.0572),
    ("2026-01-26", 2141, 299000, 299000, 0.0224, 390, 910, 900, 0.0308),
    ("2026-02-02", 2207, 325000, 325000, 0.0245, 401, 950, 950, 0.0399),
    ("2026-02-09", 3909, 315000, 315000, 0.0090, 412, 925, 925, 0.0388),
    ("2026-02-16", 3816, 325000, 325000, 0.0136, 395, 950, 950, 0.0582),
    ("2026-02-23", 4941, 309000, 309400, 0.0061, 395, 950, 950, 0.0405),
    ("2026-03-02", 1968, 300000, 300000, 0.0320, 435, 920, 910, 0.0368),
    ("2026-03-09", 959, 315000, 315000, 0.0448, 411, 950, 950, 0.0438),
    ("2026-03-16", 1056, 299000, 299000, 0.0417, 419, 950, 950, 0.0358),
    ("2026-03-23", 1040, 315000, 314625, 0.0346, 364, 925, 925, 0.0440),
    ("2026-03-30", 1286, 334800, 334400, 0.0280, 411, 950, 950, 0.0365),
    ("2026-04-06", 767, 302500, 302500, 0.0326, 352, 945, 940, 0.0284),
    ("2026-04-13", 898, 296800, 296800, 0.0256, 364, 900, 900, 0.0357),
    ("2026-04-20", 1075, 322000, 322000, 0.0251, 447, 950, 950, 0.0246),
    ("2026-04-27", 858, 315000, 315000, 0.0233, 332, 925, 925, 0.0241),
    ("2026-05-04", 1083, 322000, 320000, 0.0212, 447, 950, 950, 0.0112),
    ("2026-05-11", 778, 321250, 319750, 0.0103, 322, 900, 900, 0.0311),
    ("2026-05-18", 1233, 334900, 334900, 0.0146, 430, 950, 950, 0.0093),
    ("2026-05-25", 882, 321000, 321000, 0.0034, 345, 950, 950, 0.0029),
    ("2026-06-01", 1076, 316273, 317500, 0.0056, 395, 950, 950, 0.0025),
    ("2026-06-08", 565, 329500, 329500, 0.0, 177, 950, 950, 0.0),
]

HEADER_FILL = PatternFill("solid", fgColor="385723")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NUM = "#,##0"
EUR = "€ #,##0"
PCT = "0.0%"
DAYS = "#,##0.0"


def weeklabel(datestr):
    d = datetime.date.fromisoformat(datestr)
    iso = d.isocalendar()
    return f"{iso[0]}-W{iso[1]:02d}"


def make_sheet(wb, title, headers, rows, numfmts=None, widths=None):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if numfmts and c in numfmts:
                cell.number_format = numfmts[c]
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 12)
    ws.freeze_panes = "A2"
    return ws


wb = Workbook()
wb.remove(wb.active)

# --- Op de markt gekomen (objecten) ---
rows = []
for w, tot, eerste, gap, pubs in MARKT:
    her = tot - eerste
    rows.append((weeklabel(w), datetime.date.fromisoformat(w), tot, eerste, her,
                 her / tot if tot else None, gap, pubs))
make_sheet(
    wb, "Op de markt gekomen",
    ["Week", "Maandag", "Unieke objecten op de markt", "Waarvan eerste keer",
     "Waarvan herpublicatie", "% herpublicatie", "Mediaan dagen sinds vorige periode",
     "Nieuwe publicaties (ter vergelijking)"],
    rows, numfmts={3: NUM, 4: NUM, 5: NUM, 6: PCT, 7: DAYS, 8: NUM},
    widths={2: 12, 3: 15, 4: 13, 5: 13, 6: 12, 7: 16, 8: 16},
)

# --- Doorlooptijd (van de markt gegaan) ---
rows = [(weeklabel(w), datetime.date.fromisoformat(w), kn, kd, hn, hd)
        for w, kn, kd, hn, hd in DOORLOOPTIJD]
make_sheet(
    wb, "Doorlooptijd",
    ["Week", "Maandag", "Koop: objecten van de markt", "Koop: mediaan dagen op de markt",
     "Huur: objecten van de markt", "Huur: mediaan dagen op de markt"],
    rows, numfmts={3: NUM, 4: DAYS, 5: NUM, 6: DAYS},
    widths={2: 12, 3: 15, 4: 16, 5: 15, 6: 16},
)

# --- Prijsverlagingen ---
rows = [(weeklabel(w), datetime.date.fromisoformat(w), kn, ke, kp / 100.0, hn, he, hp / 100.0)
        for w, kn, ke, kp, hn, he, hp in PRIJSVERLAGINGEN]
make_sheet(
    wb, "Prijsverlagingen",
    ["Week", "Maandag", "Koop: aantal verlagingen", "Koop: mediaan daling",
     "Koop: mediaan daling %", "Huur: aantal verlagingen", "Huur: mediaan daling (per mnd)",
     "Huur: mediaan daling %"],
    rows, numfmts={3: NUM, 4: EUR, 5: PCT, 6: NUM, 7: EUR, 8: PCT},
    widths={2: 12, 3: 14, 4: 14, 5: 13, 6: 14, 7: 16, 8: 13},
)

# --- Startprijs vs huidige prijs (huis / appartement) ---
STARTKOLOMMEN = ["Week", "Maandag",
                 "Koop N", "Koop med. startprijs", "Koop med. huidige prijs", "Koop % verlaagd",
                 "Huur N", "Huur med. startprijs", "Huur med. huidige prijs", "Huur % verlaagd"]
STARTFMT = {3: NUM, 4: EUR, 5: EUR, 6: PCT, 7: NUM, 8: EUR, 9: EUR, 10: PCT}
STARTBREEDTE = {2: 12, 4: 15, 5: 16, 6: 12, 8: 15, 9: 16, 10: 12}


def startrows(data):
    return [(weeklabel(w), datetime.date.fromisoformat(w), kn, ks, kl, kv, hn, hs, hl, hv)
            for w, kn, ks, kl, kv, hn, hs, hl, hv in data]


make_sheet(wb, "Startprijs huis", STARTKOLOMMEN, startrows(START_HUIS),
           numfmts=STARTFMT, widths=STARTBREEDTE)
make_sheet(wb, "Startprijs appartement", STARTKOLOMMEN, startrows(START_APP),
           numfmts=STARTFMT, widths=STARTBREEDTE)

out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, "spotto-weekrapport-2026-sectie2-objecten.xlsx")
wb.save(out_path)
print(f"OK: {out_path}")
