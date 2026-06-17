# Bouwt het wekelijkse Spotto-rapport, sectie 3 (Bezoekers & zoekers - GA4).
# Bron: Google Analytics 4, property "Spotto - V2" (491908260), opgehaald 2026-06-12.
# Weken = ISO-weken (dimensie isoYearIsoWeek, tijdzone Europe/Brussels).
# Week 24 is de lopende week (t/m 12 juni).
import datetime
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# (isoweek, totale bezoekers, nieuwe bezoekers, sessies, engagement rate)
BEZOEKERS = [
    ("202601", 34439, 24013, 55846, 0.676), ("202602", 39866, 27064, 69416, 0.681),
    ("202603", 42884, 29152, 74434, 0.697), ("202604", 41110, 27938, 71032, 0.688),
    ("202605", 36718, 23735, 65264, 0.693), ("202606", 40755, 27804, 71712, 0.687),
    ("202607", 43234, 28594, 75192, 0.714), ("202608", 44395, 29986, 78644, 0.746),
    ("202609", 41282, 27668, 72471, 0.725), ("202610", 40084, 26249, 68849, 0.709),
    ("202611", 41222, 27122, 71476, 0.717), ("202612", 39429, 25636, 68728, 0.697),
    ("202613", 40649, 26162, 71173, 0.714), ("202614", 37776, 24234, 65932, 0.688),
    ("202615", 34158, 22139, 58906, 0.707), ("202616", 37507, 24645, 65827, 0.716),
    ("202617", 36688, 24110, 65625, 0.716), ("202618", 36050, 23994, 64475, 0.701),
    ("202619", 37297, 24916, 66609, 0.691), ("202620", 38147, 25373, 66866, 0.727),
    ("202621", 34051, 22475, 62384, 0.701), ("202622", 31971, 20841, 55573, 0.693),
    ("202623", 35099, 23715, 63951, 0.691), ("202624", 23601, 14491, 38750, 0.626),
]

# (isoweek, zoekers koop, zoekers huur) - unieke gebruikers die >=1 pandpagina (overzicht of
# detail) met dat transactietype bekeken; overlap mogelijk, dus niet optellen
ZOEKERS = [
    ("202601", 21722, 11953), ("202602", 23916, 14731), ("202603", 26489, 15607),
    ("202604", 24961, 15193), ("202605", 22561, 13587), ("202606", 25086, 14888),
    ("202607", 26176, 16028), ("202608", 26993, 16523), ("202609", 25066, 15597),
    ("202610", 23826, 15400), ("202611", 24667, 15809), ("202612", 24076, 14860),
    ("202613", 24436, 15055), ("202614", 23056, 13853), ("202615", 21134, 12102),
    ("202616", 23451, 13125), ("202617", 22904, 12937), ("202618", 22865, 12449),
    ("202619", 23423, 12851), ("202620", 24473, 12762), ("202621", 21492, 12224),
    ("202622", 20028, 10991), ("202623", 22113, 12502), ("202624", 14802, 7771),
]

# (isoweek, overzicht koop, overzicht huur, detail koop, detail huur)
PANDWEERGAVES = [
    ("202601", 134556, 74736, 31016, 15513), ("202602", 145491, 93005, 34830, 20265),
    ("202603", 159974, 99003, 38568, 21423), ("202604", 159266, 97505, 37931, 21590),
    ("202605", 141798, 87631, 32648, 18416), ("202606", 161342, 96943, 37215, 20631),
    ("202607", 168094, 107128, 37492, 23175), ("202608", 165879, 105631, 38507, 22315),
    ("202609", 147371, 103876, 33368, 20596), ("202610", 139160, 101662, 30396, 19645),
    ("202611", 147428, 103094, 34607, 20033), ("202612", 137475, 95583, 31404, 18938),
    ("202613", 141575, 99171, 32417, 19710), ("202614", 131610, 89137, 30411, 17836),
    ("202615", 120247, 81742, 27732, 16704), ("202616", 129582, 83841, 31533, 17082),
    ("202617", 128658, 88892, 30663, 18477), ("202618", 133921, 82605, 30482, 16890),
    ("202619", 132463, 84277, 31512, 17689), ("202620", 136826, 88604, 31869, 18448),
    ("202621", 117718, 82712, 28232, 17418), ("202622", 112371, 70702, 26911, 14374),
    ("202623", 123534, 82190, 29342, 17786), ("202624", 81723, 45010, 18747, 9758),
]

# (isoweek, zoekopdrachten, unieke zoekende gebruikers) - event property_search
ZOEKOPDRACHTEN = [
    ("202601", 7300, 2890), ("202602", 8676, 3246), ("202603", 9338, 3393),
    ("202604", 9099, 3502), ("202605", 8445, 3153), ("202606", 9529, 3614),
    ("202607", 9658, 3767), ("202608", 10543, 3940), ("202609", 8680, 3470),
    ("202610", 8637, 3414), ("202611", 9240, 3543), ("202612", 8421, 3360),
    ("202613", 9374, 3492), ("202614", 8336, 3188), ("202615", 7386, 2951),
    ("202616", 7691, 3029), ("202617", 8109, 3086), ("202618", 7717, 3036),
    ("202619", 8214, 3245), ("202620", 7875, 3094), ("202621", 7023, 2990),
    ("202622", 7169, 2750), ("202623", 7754, 3019), ("202624", 4575, 1892),
]

# (isoweek, Organic Search, Paid Search, Direct, Unassigned, Email, Cross-network,
#  Referral, Organic Social, Overig) - sessies per kanaal
KANALEN = [
    ("202601", 31223, 17304, 3989, 1094, 591, 450, 531, 393, 260),
    ("202602", 40142, 18088, 5139, 1370, 1197, 455, 662, 383, 302),
    ("202603", 45305, 18740, 5913, 1944, 1039, 473, 651, 367, 281),
    ("202604", 43034, 16978, 5544, 1630, 1051, 436, 618, 429, 263),
    ("202605", 41301, 13964, 4787, 1162, 946, 341, 707, 342, 264),
    ("202606", 44349, 16865, 5106, 1495, 1077, 367, 754, 358, 336),
    ("202607", 46546, 18053, 5653, 1471, 1290, 664, 700, 454, 324),
    ("202608", 52261, 16147, 5640, 1303, 1253, 818, 737, 440, 366),
    ("202609", 50082, 13186, 5418, 1316, 1185, 594, 797, 406, 53),
    ("202610", 47403, 12729, 4904, 1373, 1149, 581, 770, 268, 20),
    ("202611", 49098, 13098, 5015, 1377, 1266, 627, 730, 297, 21),
    ("202612", 49072, 11605, 4854, 1250, 1197, 567, 711, 400, 20),
    ("202613", 49144, 12443, 5048, 1439, 1317, 550, 654, 525, 25),
    ("202614", 44750, 11850, 4926, 1588, 1256, 592, 537, 364, 15),
    ("202615", 39163, 11671, 4324, 1322, 1228, 583, 543, 368, 10),
    ("202616", 43538, 12071, 4694, 1237, 1354, 547, 602, 1026, 12),
    ("202617", 44752, 11422, 5245, 1257, 1428, 535, 605, 399, 6),
    ("202618", 43767, 11983, 4908, 1079, 1294, 594, 615, 342, 11),
    ("202619", 44647, 12271, 4911, 1228, 1454, 580, 725, 493, 9),
    ("202620", 44762, 12604, 4821, 1505, 1235, 687, 663, 659, 17),
    ("202621", 39508, 12375, 4883, 1162, 1834, 632, 687, 403, 7),
    ("202622", 35966, 11311, 4584, 1168, 1228, 509, 678, 237, 14),
    ("202623", 41012, 12258, 5201, 1286, 1565, 607, 723, 668, 7),
    # W24 hertrokken op 12/6 namiddag: kanaalattributie van de lopende dag is instabiel
    # (betaalde sessies staan tijdelijk als Unassigned/Cross-network tot Google Ads-koppeling
    # verwerkt is). Definitieve cijfers pas na afsluiten van de week.
    ("202624", 24490, 8095, 3122, 2026, 899, 1095, 513, 295, 110),
]

HEADER_FILL = PatternFill("solid", fgColor="BF5B17")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NUM = "#,##0"
PCT = "0.0%"


def weekinfo(isoweek):
    year, wk = int(isoweek[:4]), int(isoweek[4:])
    monday = datetime.date.fromisocalendar(year, wk, 1)
    return f"{year}-W{wk:02d}", monday


def make_sheet(wb, title, headers, rows, numfmts=None, widths=None):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if numfmts and c in numfmts:
                cell.number_format = numfmts[c]
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 12)
    ws.freeze_panes = "A2"
    return ws


wb = Workbook()
wb.remove(wb.active)

# --- Bezoekers ---
rows = [(*weekinfo(w), tot, nieuw, ses, rate) for w, tot, nieuw, ses, rate in BEZOEKERS]
make_sheet(
    wb, "Bezoekers",
    ["Week", "Maandag", "Bezoekers", "Nieuwe bezoekers", "Sessies", "Engagement rate"],
    rows, numfmts={3: NUM, 4: NUM, 5: NUM, 6: PCT},
    widths={2: 12, 4: 14, 6: 13},
)

# --- Zoekers koop vs huur ---
bez = {w: tot for w, tot, _, _, _ in BEZOEKERS}
rows = [(*weekinfo(w), k, h, bez[w]) for w, k, h in ZOEKERS]
make_sheet(
    wb, "Zoekers koop vs huur",
    ["Week", "Maandag", "Zoekers te koop", "Zoekers te huur", "Totale bezoekers (referentie)"],
    rows, numfmts={3: NUM, 4: NUM, 5: NUM},
    widths={2: 12, 3: 13, 4: 13, 5: 16},
)

# --- Pandweergaves ---
rows = [(*weekinfo(w), ok, oh, dk, dh, ok + oh + dk + dh)
        for w, ok, oh, dk, dh in PANDWEERGAVES]
make_sheet(
    wb, "Pandweergaves",
    ["Week", "Maandag", "Overzichtspagina koop", "Overzichtspagina huur",
     "Detailpagina koop", "Detailpagina huur", "Totaal"],
    rows, numfmts={c: NUM for c in range(3, 8)},
    widths={2: 12, 3: 14, 4: 14, 5: 13, 6: 13},
)

# --- Zoekopdrachten ---
rows = [(*weekinfo(w), n, u) for w, n, u in ZOEKOPDRACHTEN]
make_sheet(
    wb, "Zoekopdrachten",
    ["Week", "Maandag", "Zoekopdrachten", "Unieke zoekende gebruikers"],
    rows, numfmts={3: NUM, 4: NUM},
    widths={2: 12, 3: 14, 4: 16},
)

# --- Verkeersbronnen ---
rows = []
for w, *vals in KANALEN:
    rows.append((*weekinfo(w), *vals, sum(vals)))
make_sheet(
    wb, "Verkeersbronnen",
    ["Week", "Maandag", "Organic Search", "Paid Search", "Direct", "Unassigned",
     "Email", "Cross-network", "Referral", "Organic Social", "Overig", "Totaal sessies"],
    rows, numfmts={c: NUM for c in range(3, 13)},
    widths={2: 12, 3: 13, 8: 13, 10: 13, 12: 13},
)

out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, "spotto-weekrapport-2026-sectie3-bezoekers.xlsx")
wb.save(out_path)
print(f"OK: {out_path}")
