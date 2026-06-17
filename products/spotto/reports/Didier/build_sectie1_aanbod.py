# Bouwt het wekelijkse Spotto-rapport, sectie 1 (Aanbod - publicaties).
# Data komt uit de Spotto-productiedatabank (dbo.Publications), opgehaald 2026-06-12.
# Weken = ISO-weken (maandag t/m zondag). Week 24 is de lopende week (t/m 12 juni).
# Telcijfers gesplitst op pandtype: huis / appartement / overig.
# Vraagprijzen en EPC apart voor huis en appartement.
import datetime
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# (weekstart, K huis, K app, K overig, H huis, H app, H overig) - nieuw gepubliceerd (Online_From)
NIEUW = [
    ("2025-12-29", 441, 362, 265, 52, 147, 380),
    ("2026-01-05", 1459, 1180, 446, 152, 473, 302),
    ("2026-01-12", 1143, 1090, 432, 137, 440, 277),
    ("2026-01-19", 1514, 1408, 433, 154, 437, 245),
    ("2026-01-26", 2019, 2156, 780, 115, 390, 248),
    ("2026-02-02", 2126, 2212, 665, 150, 401, 262),
    ("2026-02-09", 3425, 3911, 937, 127, 412, 177),
    ("2026-02-16", 3778, 3838, 1009, 142, 395, 243),
    ("2026-02-23", 4518, 4945, 1068, 108, 395, 218),
    ("2026-03-02", 1911, 1998, 622, 153, 435, 266),
    ("2026-03-09", 1276, 987, 406, 94, 411, 253),
    ("2026-03-16", 1333, 1156, 441, 134, 419, 279),
    ("2026-03-23", 1583, 1074, 459, 132, 364, 262),
    ("2026-03-30", 1393, 1306, 478, 132, 411, 286),
    ("2026-04-06", 978, 779, 288, 111, 352, 177),
    ("2026-04-13", 1232, 957, 434, 116, 365, 212),
    ("2026-04-20", 1554, 1175, 610, 145, 447, 262),
    ("2026-04-27", 1202, 870, 391, 127, 332, 200),
    ("2026-05-04", 1591, 1150, 528, 143, 447, 276),
    ("2026-05-11", 1219, 787, 362, 104, 322, 191),
    ("2026-05-18", 1580, 1268, 629, 135, 430, 261),
    ("2026-05-25", 1281, 934, 369, 122, 345, 142),
    ("2026-06-01", 1467, 1104, 526, 143, 395, 1202),
    ("2026-06-08", 1080, 1032, 350, 101, 347, 169),
]

# (weekstart, K huis, K app, K overig, H huis, H app, H overig) - offline gehaald (Online_To)
OFFLINE = [
    ("2025-12-29", 758, 683, 213, 67, 177, 113),
    ("2026-01-05", 1445, 1240, 505, 113, 358, 295),
    ("2026-01-12", 1627, 1345, 647, 136, 441, 236),
    ("2026-01-19", 1592, 1531, 416, 142, 401, 227),
    ("2026-01-26", 2047, 2213, 733, 138, 426, 247),
    ("2026-02-02", 2163, 2242, 661, 144, 391, 212),
    ("2026-02-09", 3521, 3832, 926, 161, 409, 194),
    ("2026-02-16", 3753, 3718, 956, 149, 407, 232),
    ("2026-02-23", 4665, 4916, 1129, 116, 433, 231),
    ("2026-03-02", 1859, 1836, 550, 124, 377, 200),
    ("2026-03-09", 1148, 1078, 443, 133, 430, 187),
    ("2026-03-16", 1312, 1148, 461, 119, 401, 425),
    ("2026-03-23", 1412, 1047, 513, 142, 395, 235),
    ("2026-03-30", 1459, 1390, 598, 138, 442, 299),
    ("2026-04-06", 868, 798, 258, 94, 323, 176),
    ("2026-04-13", 1125, 911, 392, 119, 372, 197),
    ("2026-04-20", 1495, 1186, 478, 153, 417, 230),
    ("2026-04-27", 1111, 919, 370, 120, 414, 194),
    ("2026-05-04", 1449, 958, 485, 130, 365, 263),
    ("2026-05-11", 1074, 761, 322, 100, 310, 200),
    ("2026-05-18", 1610, 1464, 559, 128, 475, 260),
    ("2026-05-25", 1206, 938, 386, 120, 371, 187),
    ("2026-06-01", 1474, 1061, 407, 140, 402, 285),
    ("2026-06-08", 1066, 916, 348, 106, 319, 157),
]

# (peildatum, K huis, K app, K ov, H huis, H app, H ov, kantoren koop, huur, totaal)
STOCK = [
    ("2026-01-05", 13804, 14265, 8922, 647, 2071, 5568, 903, 675, 937),
    ("2026-01-12", 13818, 14205, 8863, 686, 2186, 5575, 906, 679, 938),
    ("2026-01-19", 13334, 13950, 8648, 687, 2185, 5616, 889, 661, 918),
    ("2026-01-26", 13257, 13826, 8666, 699, 2221, 5634, 878, 667, 907),
    ("2026-02-02", 13235, 13773, 8712, 676, 2185, 5635, 879, 664, 907),
    ("2026-02-09", 13177, 13725, 8720, 682, 2195, 5685, 877, 661, 905),
    ("2026-02-16", 13068, 13810, 8732, 648, 2198, 5668, 879, 661, 909),
    ("2026-02-23", 13097, 13932, 8783, 641, 2186, 5679, 880, 669, 914),
    ("2026-03-02", 12990, 13976, 8719, 633, 2148, 5666, 884, 655, 915),
    ("2026-03-09", 13025, 14130, 8791, 662, 2206, 5732, 885, 656, 916),
    ("2026-03-16", 13153, 14039, 8754, 623, 2187, 5798, 889, 660, 919),
    ("2026-03-23", 13174, 14047, 8734, 638, 2205, 5652, 897, 658, 929),
    ("2026-03-30", 13345, 14074, 8680, 628, 2174, 5679, 894, 658, 927),
    ("2026-04-06", 13279, 13990, 8560, 622, 2143, 5666, 882, 649, 911),
    ("2026-04-13", 13389, 13971, 8590, 639, 2172, 5667, 879, 649, 909),
    ("2026-04-20", 13496, 14017, 8632, 636, 2165, 5682, 878, 648, 905),
    ("2026-04-27", 13555, 14006, 8764, 628, 2195, 5714, 879, 633, 907),
    ("2026-05-04", 13646, 13957, 8785, 635, 2113, 5720, 890, 639, 917),
    ("2026-05-11", 13788, 14149, 8828, 648, 2195, 5733, 893, 641, 920),
    ("2026-05-18", 13933, 14175, 8868, 652, 2207, 5724, 891, 644, 917),
    ("2026-05-25", 13903, 13979, 8938, 659, 2162, 5725, 898, 645, 921),
    ("2026-06-01", 13978, 13975, 8921, 661, 2136, 5680, 907, 656, 928),
    ("2026-06-08", 13971, 14018, 9040, 664, 2129, 6597, 904, 654, 928),
    ("2026-06-12", 13992, 14142, 9054, 656, 2162, 6613, 907, 661, 929),
]

# (weekstart, koop mediaan, koop gem, koop N, huur mediaan, huur gem, huur N) - HUIZEN
PRIJZEN_HUIS = [
    ("2025-12-29", 395000, 481619, 429, 1100, 1184, 50),
    ("2026-01-05", 419000, 505783, 1399, 1200, 1370, 149),
    ("2026-01-12", 415000, 480682, 1090, 1195, 1317, 135),
    ("2026-01-19", 395000, 462021, 1425, 1195, 1281, 153),
    ("2026-01-26", 399000, 451598, 1875, 1100, 1332, 111),
    ("2026-02-02", 425000, 515170, 2012, 1100, 1263, 147),
    ("2026-02-09", 439000, 478595, 3297, 1125, 1264, 127),
    ("2026-02-16", 439000, 512983, 3436, 1250, 1405, 139),
    ("2026-02-23", 447000, 520218, 4232, 1250, 1298, 108),
    ("2026-03-02", 415000, 499554, 1827, 1095, 1259, 152),
    ("2026-03-09", 425000, 485095, 1198, 1100, 1313, 94),
    ("2026-03-16", 417708, 493448, 1240, 1225, 1309, 130),
    ("2026-03-23", 430714, 507203, 1471, 1200, 1293, 130),
    ("2026-03-30", 425000, 524636, 1325, 1200, 1381, 131),
    ("2026-04-06", 420000, 507190, 933, 1150, 1326, 110),
    ("2026-04-13", 399000, 475283, 1184, 1150, 1331, 115),
    ("2026-04-20", 429000, 506644, 1483, 1250, 1438, 143),
    ("2026-04-27", 419000, 486699, 1144, 1200, 1353, 126),
    ("2026-05-04", 419000, 502850, 1545, 1200, 1314, 143),
    ("2026-05-11", 435000, 508725, 1176, 1200, 1360, 103),
    ("2026-05-18", 429000, 511404, 1514, 1200, 1372, 135),
    ("2026-05-25", 429000, 503894, 1221, 1195, 1320, 121),
    ("2026-06-01", 422905, 520129, 1432, 1250, 1395, 143),
    ("2026-06-08", 425000, 479750, 1058, 1150, 1369, 101),
]

# idem - APPARTEMENTEN
PRIJZEN_APP = [
    ("2025-12-29", 299000, 417538, 357, 950, 1049, 144),
    ("2026-01-05", 299000, 374407, 1144, 900, 998, 467),
    ("2026-01-12", 299500, 343052, 1041, 900, 1013, 437),
    ("2026-01-19", 299000, 342429, 1384, 940, 1022, 433),
    ("2026-01-26", 299000, 342052, 2117, 900, 1001, 390),
    ("2026-02-02", 325000, 402548, 2173, 950, 1022, 396),
    ("2026-02-09", 315000, 376349, 3824, 925, 1009, 405),
    ("2026-02-16", 325000, 395185, 3743, 950, 1055, 393),
    ("2026-02-23", 315000, 391602, 4834, 950, 1044, 385),
    ("2026-03-02", 304877, 374804, 1899, 920, 1007, 431),
    ("2026-03-09", 314600, 350965, 935, 950, 1016, 410),
    ("2026-03-16", 299000, 346724, 1037, 950, 1022, 419),
    ("2026-03-23", 310000, 360727, 1013, 925, 1007, 360),
    ("2026-03-30", 329500, 384611, 1247, 950, 1037, 409),
    ("2026-04-06", 299750, 377479, 758, 942.5, 1008, 346),
    ("2026-04-13", 295000, 377013, 884, 900, 991, 362),
    ("2026-04-20", 321000, 367437, 1059, 950, 1046, 444),
    ("2026-04-27", 315000, 371167, 840, 925, 1011, 330),
    ("2026-05-04", 317000, 368756, 1069, 950, 1019, 445),
    ("2026-05-11", 319000, 375627, 769, 900, 993, 318),
    ("2026-05-18", 332500, 407677, 1215, 950, 1044, 427),
    ("2026-05-25", 320250, 390673, 870, 950, 1062, 343),
    ("2026-06-01", 315000, 389977, 1063, 950, 1012, 395),
    ("2026-06-08", 325000, 367201, 990, 950, 1010, 344),
]

# (weekstart, K: A+,A,B,C,D,E,F,onbekend, H: idem) - HUIZEN
EPC_HUIS = [
    ("2025-12-29", 44, 85, 103, 67, 31, 62, 1, 48, 7, 11, 14, 5, 4, 0, 0, 11),
    ("2026-01-05", 175, 252, 303, 219, 90, 242, 0, 178, 38, 39, 36, 12, 5, 2, 0, 20),
    ("2026-01-12", 127, 201, 244, 142, 87, 191, 5, 146, 22, 30, 38, 17, 4, 5, 0, 21),
    ("2026-01-19", 199, 244, 350, 212, 125, 215, 1, 168, 34, 38, 38, 20, 5, 4, 0, 15),
    ("2026-01-26", 313, 309, 374, 216, 135, 301, 0, 371, 22, 26, 28, 11, 5, 4, 0, 19),
    ("2026-02-02", 351, 328, 413, 299, 124, 312, 4, 295, 27, 34, 41, 18, 4, 6, 1, 19),
    ("2026-02-09", 402, 540, 745, 549, 253, 340, 2, 594, 26, 34, 29, 12, 6, 3, 0, 17),
    ("2026-02-16", 420, 574, 872, 596, 283, 435, 2, 596, 22, 33, 37, 19, 5, 7, 0, 19),
    ("2026-02-23", 514, 572, 1120, 681, 349, 531, 2, 749, 23, 24, 26, 11, 6, 4, 0, 14),
    ("2026-03-02", 229, 287, 392, 301, 140, 289, 3, 270, 26, 39, 37, 22, 4, 4, 0, 21),
    ("2026-03-09", 201, 204, 276, 178, 83, 193, 0, 141, 23, 26, 21, 10, 1, 3, 0, 10),
    ("2026-03-16", 203, 206, 243, 175, 101, 201, 0, 204, 25, 35, 28, 14, 1, 6, 0, 25),
    ("2026-03-23", 231, 228, 229, 194, 104, 195, 4, 398, 32, 26, 32, 20, 8, 5, 0, 9),
    ("2026-03-30", 211, 230, 273, 191, 77, 206, 3, 202, 23, 39, 27, 16, 6, 7, 0, 14),
    ("2026-04-06", 157, 175, 207, 144, 72, 142, 2, 79, 23, 34, 26, 11, 5, 4, 0, 8),
    ("2026-04-13", 174, 189, 267, 170, 98, 234, 2, 98, 21, 24, 35, 17, 3, 3, 0, 13),
    ("2026-04-20", 244, 282, 305, 204, 125, 257, 3, 134, 35, 36, 40, 19, 4, 4, 0, 7),
    ("2026-04-27", 169, 220, 227, 185, 85, 192, 4, 120, 25, 28, 30, 18, 8, 2, 1, 15),
    ("2026-05-04", 233, 275, 364, 236, 98, 237, 6, 142, 26, 37, 32, 18, 7, 4, 0, 19),
    ("2026-05-11", 164, 220, 239, 175, 96, 184, 7, 134, 18, 26, 30, 16, 2, 4, 0, 8),
    ("2026-05-18", 275, 271, 304, 203, 96, 242, 4, 185, 25, 38, 30, 21, 5, 6, 0, 10),
    ("2026-05-25", 191, 230, 291, 174, 84, 184, 6, 121, 23, 22, 29, 19, 7, 6, 0, 16),
    ("2026-06-01", 190, 241, 303, 190, 89, 199, 7, 248, 31, 32, 36, 16, 5, 3, 1, 19),
    ("2026-06-08", 152, 176, 235, 136, 80, 194, 2, 105, 23, 31, 23, 7, 4, 2, 1, 10),
]

# idem - APPARTEMENTEN
EPC_APP = [
    ("2025-12-29", 62, 137, 60, 21, 3, 10, 1, 68, 40, 56, 17, 6, 2, 1, 0, 25),
    ("2026-01-05", 216, 356, 196, 63, 25, 16, 4, 304, 110, 195, 78, 17, 8, 8, 1, 56),
    ("2026-01-12", 236, 309, 143, 52, 23, 16, 5, 306, 106, 151, 80, 22, 9, 11, 2, 59),
    ("2026-01-19", 235, 408, 317, 109, 21, 12, 5, 301, 93, 148, 68, 22, 13, 7, 3, 83),
    ("2026-01-26", 467, 574, 449, 148, 16, 9, 4, 489, 99, 124, 73, 19, 8, 6, 2, 59),
    ("2026-02-02", 417, 823, 374, 159, 17, 10, 3, 409, 105, 146, 66, 19, 10, 5, 0, 50),
    ("2026-02-09", 809, 1361, 616, 244, 8, 66, 9, 798, 115, 144, 69, 25, 5, 7, 0, 47),
    ("2026-02-16", 707, 1338, 545, 220, 87, 55, 3, 883, 116, 129, 75, 20, 10, 2, 0, 43),
    ("2026-02-23", 935, 1787, 662, 272, 135, 65, 7, 1082, 107, 142, 66, 21, 6, 4, 2, 47),
    ("2026-03-02", 368, 601, 275, 88, 41, 22, 3, 600, 108, 158, 61, 24, 13, 5, 9, 57),
    ("2026-03-09", 209, 289, 144, 53, 26, 12, 3, 251, 100, 141, 66, 23, 13, 4, 2, 62),
    ("2026-03-16", 174, 369, 177, 49, 17, 14, 6, 350, 117, 143, 72, 22, 7, 8, 3, 47),
    ("2026-03-23", 238, 316, 147, 63, 17, 15, 7, 271, 76, 139, 71, 22, 7, 4, 2, 43),
    ("2026-03-30", 266, 315, 192, 54, 18, 20, 4, 437, 88, 152, 75, 25, 14, 2, 0, 55),
    ("2026-04-06", 164, 251, 137, 40, 14, 13, 1, 159, 86, 145, 58, 17, 6, 4, 3, 33),
    ("2026-04-13", 278, 262, 143, 41, 16, 10, 8, 199, 72, 136, 62, 23, 8, 10, 7, 47),
    ("2026-04-20", 239, 348, 198, 42, 12, 18, 6, 312, 127, 160, 71, 26, 8, 9, 2, 44),
    ("2026-04-27", 204, 325, 136, 37, 18, 9, 0, 141, 88, 123, 50, 20, 7, 4, 6, 34),
    ("2026-05-04", 258, 350, 162, 50, 19, 16, 12, 283, 115, 141, 74, 35, 15, 5, 0, 62),
    ("2026-05-11", 188, 305, 144, 46, 14, 13, 12, 65, 92, 105, 56, 16, 4, 6, 6, 37),
    ("2026-05-18", 296, 389, 185, 60, 29, 20, 5, 284, 117, 150, 58, 21, 5, 4, 3, 72),
    ("2026-05-25", 216, 330, 147, 42, 12, 7, 5, 175, 116, 105, 56, 15, 11, 8, 2, 32),
    ("2026-06-01", 265, 383, 192, 49, 11, 15, 8, 181, 106, 141, 61, 28, 4, 9, 2, 44),
    ("2026-06-08", 172, 239, 160, 53, 9, 20, 13, 366, 92, 117, 62, 21, 9, 8, 1, 37),
]

# (weekstart, K: bestaand, nieuwbouw, project, onbekend, H: idem) - alle pandtypes samen
NIEUWBOUW = [
    ("2025-12-29", 972, 86, 7, 3, 562, 12, 0, 5),
    ("2026-01-05", 2605, 420, 49, 11, 892, 27, 0, 8),
    ("2026-01-12", 2188, 412, 57, 8, 811, 33, 1, 9),
    ("2026-01-19", 2875, 413, 53, 14, 776, 49, 0, 11),
    ("2026-01-26", 3984, 899, 65, 7, 696, 45, 0, 12),
    ("2026-02-02", 4279, 669, 41, 14, 779, 17, 0, 17),
    ("2026-02-09", 6945, 1275, 41, 12, 686, 20, 0, 10),
    ("2026-02-16", 7124, 1432, 56, 13, 742, 23, 0, 15),
    ("2026-02-23", 8570, 1904, 46, 11, 692, 19, 0, 10),
    ("2026-03-02", 3492, 943, 64, 32, 779, 52, 0, 23),
    ("2026-03-09", 2122, 472, 58, 17, 721, 25, 0, 12),
    ("2026-03-16", 2337, 511, 71, 11, 776, 52, 0, 4),
    ("2026-03-23", 2246, 793, 65, 12, 706, 38, 1, 13),
    ("2026-03-30", 2395, 690, 75, 17, 783, 26, 1, 19),
    ("2026-04-06", 1762, 245, 30, 8, 618, 18, 0, 4),
    ("2026-04-13", 2131, 431, 46, 15, 662, 25, 0, 6),
    ("2026-04-20", 2619, 660, 47, 13, 803, 33, 2, 16),
    ("2026-04-27", 2152, 275, 25, 11, 619, 30, 1, 9),
    ("2026-05-04", 2634, 562, 42, 31, 816, 38, 1, 11),
    ("2026-05-11", 2051, 273, 36, 8, 598, 15, 0, 4),
    ("2026-05-18", 2682, 676, 101, 18, 776, 39, 0, 11),
    ("2026-05-25", 2154, 344, 63, 23, 583, 20, 0, 6),
    ("2026-06-01", 2569, 474, 45, 9, 1717, 15, 0, 8),
    ("2026-06-08", 1880, 483, 51, 5, 582, 19, 0, 6),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NUM = "#,##0"
EUR = "€ #,##0"


def weeklabel(datestr):
    d = datetime.date.fromisoformat(datestr)
    iso = d.isocalendar()
    return f"{iso[0]}-W{iso[1]:02d}"


def make_sheet(wb, title, headers, rows, numfmts=None, widths=None):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if numfmts and c in numfmts:
                cell.number_format = numfmts[c]
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 12)
    ws.freeze_panes = "A2"
    return ws


wb = Workbook()
wb.remove(wb.active)

# --- Leeswijzer (uitgeschakeld: toelichting komt in een apart document) ---
LEESWIJZER_ACTIEF = False
lines = [
    ("Spotto weekrapport 2026 — Sectie 1: Aanbod (publicaties)", True),
    ("", False),
    ("Bron: Spotto-productiedatabank (publicaties). Opgehaald op 12 juni 2026.", False),
    ("Periode: ISO-week 1 (29 dec 2025 – 4 jan 2026) t/m week 24 (lopend, cijfers t/m 12 juni).", False),
    ("Alle cijfers tellen PUBLICATIES (advertenties), geen unieke panden. Eenzelfde pand dat opnieuw", False),
    ("gepubliceerd wordt of van kantoor wisselt, telt opnieuw mee. Objectniveau volgt in sectie 2.", False),
    ("", False),
    ("Pandtype-indeling", True),
    ("Telcijfers zijn gesplitst in huis / appartement / overig. 'Overig' bundelt grond, garage,", False),
    ("commercieel vastgoed (kantoor, handelspand, horeca, industrieel), kamers en onbekend — samen", False),
    ("±18% van het nieuwe aanbod. Vraagprijzen en EPC worden enkel voor huizen en appartementen", False),
    ("gerapporteerd; voor de overige types zijn prijzen onderling niet vergelijkbaar.", False),
    ("", False),
    ("Tabbladen", True),
    ("• Nieuw gepubliceerd — publicaties die die week voor het eerst online kwamen.", False),
    ("• Online aanbod — stand van het online aanbod op het einde van elke week (zondag), plus het", False),
    ("  aantal kantoren met minstens één pand online.", False),
    ("• Offline gehaald — publicaties die die week offline gingen (verkocht/verhuurd/teruggetrokken;", False),
    ("  de reden is niet uitgesplitst).", False),
    ("• Vraagprijzen huis / appartement — mediaan en gemiddelde vraagprijs van nieuwe publicaties", False),
    ("  met een prijs (koop in €, huur in €/maand). N = aantal publicaties met prijs.", False),
    ("• EPC huis / appartement — EPC-labelverdeling van nieuwe publicaties. Labels van de drie", False),
    ("  gewesten samengevoegd op letterniveau. A++ en G kwamen niet voor.", False),
    ("• Nieuwbouw vs bestaand — alle pandtypes samen, opgesplitst: bestaand pand, nieuwbouw(unit),", False),
    ("  nieuwbouwproject (één projectpublicatie kan meerdere units omvatten), onbekend.", False),
    ("", False),
    ("Kanttekeningen bij de interpretatie", True),
    ("1. De piek in nieuw + offline in februari (koop, tot ±10.500/week) is technische churn: in- en", False),
    ("   uitstroom pieken samen terwijl het online aanbod vlak blijft. Vermoedelijk een (her)import", False),
    ("   van kantoorportefeuilles, geen marktbeweging.", False),
    ("2. Week 2026-W23 (1 juni): +1.202 nieuwe huurpublicaties in 'overig' = de commerciële portefeuille", False),
    ("   van één kantoor (Ceusters, vnl. kantoren en logistiek) die aan het platform is toegevoegd.", False),
    ("   De huurstock 'overig' steeg daardoor structureel van ±5,7k naar ±6,6k. Geen marktbeweging,", False),
    ("   en residentiële huurcijfers (huis/appartement) zijn er niet door geraakt.", False),
    ("3. Gebruik bij voorkeur de MEDIAAN voor prijsanalyse; het gemiddelde blijft gevoelig voor", False),
    ("   uitschieters. Door de typesplitsing zijn de eerdere vertekende gemiddelden wel verdwenen", False),
    ("   (die zaten in commercieel/industrieel vastgoed).", False),
    ("4. Week 2026-W01 loopt deels in 2025 (29-31 dec); week 2026-W24 is onvolledig (t/m vr 12 juni).", False),
    ("5. Transactietypes beperkt tot te koop en te huur. Marginale types (lijfrente, erfpacht, ±100", False),
    ("   publicaties) zijn weggelaten.", False),
]
if LEESWIJZER_ACTIEF:
    ws = wb.create_sheet("Leeswijzer")
    for r, (text, bold) in enumerate(lines, 1):
        cell = ws.cell(row=r, column=1, value=text)
        if bold:
            cell.font = Font(bold=True, size=12 if r == 1 else 11)
    ws.column_dimensions["A"].width = 110

# --- Nieuw gepubliceerd / Offline gehaald (zelfde lay-out) ---
TELKOLOMMEN = [
    "Week", "Maandag",
    "Koop huis", "Koop app.", "Koop overig", "Koop totaal",
    "Huur huis", "Huur app.", "Huur overig", "Huur totaal", "Totaal",
]


def telrows(data):
    rows = []
    for w, kh, ka, ko, hh, ha, ho in data:
        rows.append((weeklabel(w), datetime.date.fromisoformat(w),
                     kh, ka, ko, kh + ka + ko, hh, ha, ho, hh + ha + ho,
                     kh + ka + ko + hh + ha + ho))
    return rows


make_sheet(wb, "Nieuw gepubliceerd", TELKOLOMMEN, telrows(NIEUW),
           numfmts={c: NUM for c in range(3, 12)}, widths={2: 12})
make_sheet(wb, "Offline gehaald", TELKOLOMMEN, telrows(OFFLINE),
           numfmts={c: NUM for c in range(3, 12)}, widths={2: 12})

# --- Online aanbod ---
rows = []
for snap, kh, ka, ko, hh, ha, ho, kk, kht, kt in STOCK:
    snap_d = datetime.date.fromisoformat(snap)
    # peildatum maandag 00:00 = stand op het einde van de voorbije week (zondag)
    iso = (snap_d - datetime.timedelta(days=1)).isocalendar() if snap_d.isoweekday() == 1 else snap_d.isocalendar()
    rows.append((f"{iso[0]}-W{iso[1]:02d}", snap_d,
                 kh, ka, ko, kh + ka + ko, hh, ha, ho, hh + ha + ho,
                 kh + ka + ko + hh + ha + ho, kk, kht, kt))
make_sheet(
    wb, "Online aanbod",
    ["Week", "Peildatum",
     "Koop huis", "Koop app.", "Koop overig", "Koop totaal",
     "Huur huis", "Huur app.", "Huur overig", "Huur totaal", "Totaal online",
     "Kantoren met koopaanbod", "Kantoren met huuraanbod", "Kantoren totaal"],
    rows, numfmts={c: NUM for c in range(3, 15)},
    widths={2: 12, 11: 13, 12: 16, 13: 16, 14: 13},
)

# --- Vraagprijzen huis / appartement ---
PRIJSKOLOMMEN = ["Week", "Maandag", "Koop mediaan", "Koop gemiddelde", "Koop N",
                 "Huur mediaan (per mnd)", "Huur gemiddelde (per mnd)", "Huur N"]
PRIJSFMT = {3: EUR, 4: EUR, 5: NUM, 6: EUR, 7: EUR, 8: NUM}
PRIJSBREEDTE = {2: 12, 3: 14, 4: 16, 6: 18, 7: 20}


def prijsrows(data):
    return [(weeklabel(w), datetime.date.fromisoformat(w), km, kg, kn, hm, hg, hn)
            for w, km, kg, kn, hm, hg, hn in data]


make_sheet(wb, "Vraagprijzen huis", PRIJSKOLOMMEN, prijsrows(PRIJZEN_HUIS),
           numfmts=PRIJSFMT, widths=PRIJSBREEDTE)
make_sheet(wb, "Vraagprijzen appartement", PRIJSKOLOMMEN, prijsrows(PRIJZEN_APP),
           numfmts=PRIJSFMT, widths=PRIJSBREEDTE)

# --- EPC huis / appartement ---
EPCKOLOMMEN = ["Week", "Maandag",
               "Koop A+", "Koop A", "Koop B", "Koop C", "Koop D", "Koop E", "Koop F", "Koop onbekend",
               "Huur A+", "Huur A", "Huur B", "Huur C", "Huur D", "Huur E", "Huur F", "Huur onbekend"]
EPCFMT = {c: NUM for c in range(3, 19)}
EPCBREEDTE = {2: 12, **{c: 9 for c in range(3, 19)}, 10: 13, 18: 13}


def epcrows(data):
    return [(weeklabel(r[0]), datetime.date.fromisoformat(r[0]), *r[1:]) for r in data]


make_sheet(wb, "EPC huis", EPCKOLOMMEN, epcrows(EPC_HUIS), numfmts=EPCFMT, widths=EPCBREEDTE)
make_sheet(wb, "EPC appartement", EPCKOLOMMEN, epcrows(EPC_APP), numfmts=EPCFMT, widths=EPCBREEDTE)

# --- Nieuwbouw vs bestaand ---
rows = []
for w, kb, kn, kp, ko, hb, hn, hp, ho in NIEUWBOUW:
    ktot = kb + kn + kp + ko
    htot = hb + hn + hp + ho
    rows.append((weeklabel(w), datetime.date.fromisoformat(w),
                 kb, kn, kp, ko, (kn + kp) / ktot if ktot else None,
                 hb, hn, hp, ho, (hn + hp) / htot if htot else None))
make_sheet(
    wb, "Nieuwbouw vs bestaand",
    ["Week", "Maandag",
     "Koop bestaand", "Koop nieuwbouw", "Koop project", "Koop onbekend", "Koop % nieuwbouw",
     "Huur bestaand", "Huur nieuwbouw", "Huur project", "Huur onbekend", "Huur % nieuwbouw"],
    rows,
    numfmts={3: NUM, 4: NUM, 5: NUM, 6: NUM, 7: "0.0%", 8: NUM, 9: NUM, 10: NUM, 11: NUM, 12: "0.0%"},
    widths={2: 12, 7: 14, 12: 14},
)

out_dir = os.path.dirname(os.path.abspath(__file__))
out_path = os.path.join(out_dir, "spotto-weekrapport-2026-sectie1-aanbod.xlsx")
wb.save(out_path)
print(f"OK: {out_path}")
