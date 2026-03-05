import argparse
import os
import pandas as pd
import numpy as np
import warnings
warnings.filterwarnings("ignore")

parser = argparse.ArgumentParser(description="Mediaan- en gemiddelde prijs per EPC-groep per gemeente")
parser.add_argument("--excel", type=str, nargs="?", const="analyse_prijs_per_epc_groep.xlsx", default=None,
                    help="Export to Excel. Optionally specify filename (default: results/analyse_prijs_per_epc_groep.xlsx)")
parser.add_argument("--jaar", type=int, default=None,
                    help="Filter op een specifiek jaar (bv. --jaar 2025). Default: alle jaren.")
args = parser.parse_args()

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(SCRIPT_DIR, "source", "data_kustinventaris.csv")
RESULTS_DIR = os.path.join(SCRIPT_DIR, "results")

df = pd.read_csv(CSV, na_values=["null", "NULL", "Null", ""], encoding="latin-1", low_memory=False)

# Filter: residential, sale, valid price & EPC
df = df[df["is_residential"] == True].copy()
df = df[df["transaction_type"] == "Sale"].copy()
df["price_value"] = pd.to_numeric(df["price_value"], errors="coerce")
df = df[df["price_value"] > 50000]
df = df[df["epc_label_display"].notna()]

# Parse year from online_from
df["online_from"] = pd.to_datetime(df["online_from"], errors="coerce", dayfirst=True)
df["jaar"] = df["online_from"].dt.year

# EPC grouping (A-D includes A+, A++)
good_epc = ["A", "A+", "A++", "B", "C", "D"]
poor_epc = ["E", "F"]
df["epc_groep"] = df["epc_label_display"].apply(
    lambda x: "A-D" if x in good_epc else ("E/F" if x in poor_epc else None)
)
df = df[df["epc_groep"].notna()]

print("Records na filtering: {:,}".format(len(df)))

# Apply year filter if specified
if args.jaar:
    df = df[df["jaar"] == args.jaar]
    print("Gefilterd op jaar {}: {:,} records".format(args.jaar, len(df)))
print()


# ============================================================
# Table builder: per gemeente, mediaan + gemiddelde
# ============================================================
def build_table(data, property_type):
    d = data[data["property_type"] == property_type].copy()

    # Mediaan per gemeente per EPC-groep
    med = d.groupby(["municipality_name", "epc_groep"])["price_value"].median().unstack()
    med.columns = ["Mediaan A-D", "Mediaan E/F"]

    # Gemiddelde per gemeente per EPC-groep
    avg = d.groupby(["municipality_name", "epc_groep"])["price_value"].mean().unstack()
    avg.columns = ["Gem. prijs A-D", "Gem. prijs E/F"]

    # Counts per gemeente per EPC-groep
    cnt = d.groupby(["municipality_name", "epc_groep"])["price_value"].count().unstack()
    cnt.columns = ["n A-D", "n E/F"]

    # Totaal per gemeente
    total = d.groupby("municipality_name")["price_value"].count()
    total.name = "n totaal"

    # Merge
    result = med.join(avg).join(cnt).join(total)
    result["Verschil mediaan"] = result["Mediaan E/F"] - result["Mediaan A-D"]
    result["Verschil mediaan %"] = (result["Verschil mediaan"] / result["Mediaan A-D"]) * 100
    result["% E/F van totaal"] = (result["n E/F"] / result["n totaal"]) * 100

    result = result.sort_values("% E/F van totaal", ascending=False)

    result = result[[
        "Mediaan A-D", "Mediaan E/F", "Verschil mediaan", "Verschil mediaan %",
        "Gem. prijs A-D", "Gem. prijs E/F",
        "% E/F van totaal", "n A-D", "n E/F", "n totaal"
    ]]

    return result


def _summary_row(d, label):
    d_good = d[d["epc_groep"] == "A-D"]["price_value"]
    d_poor = d[d["epc_groep"] == "E/F"]["price_value"]
    n_good = len(d_good)
    n_poor = len(d_poor)
    n_total = n_good + n_poor

    med_good = d_good.median()
    med_poor = d_poor.median()
    diff_med = med_poor - med_good
    diff_med_pct = (diff_med / med_good) * 100

    return pd.DataFrame([{
        "Mediaan A-D": med_good,
        "Mediaan E/F": med_poor,
        "Verschil mediaan": diff_med,
        "Verschil mediaan %": diff_med_pct,
        "Gem. prijs A-D": d_good.mean(),
        "Gem. prijs E/F": d_poor.mean(),
        "% E/F van totaal": (n_poor / n_total) * 100,
        "n A-D": n_good,
        "n E/F": n_poor,
        "n totaal": n_total,
    }], index=[label])


def add_total_row(tbl, data, property_type):
    d = data[data["property_type"] == property_type]
    total = _summary_row(d, "KUSTBREED")
    d_excl = d[d["municipality_name"] != "Knokke-Heist"]
    excl = _summary_row(d_excl, "KUST EXCL. KNOKKE")
    return pd.concat([tbl, total, excl])


# Build main tables
tbl_houses = add_total_row(build_table(df, "House"), df, "House")
tbl_apartments = add_total_row(build_table(df, "Apartment"), df, "Apartment")


def print_table(tbl, label):
    print("=" * 140)
    print(label)
    print("=" * 140)
    display = tbl.copy()
    for col in ["Mediaan A-D", "Mediaan E/F", "Gem. prijs A-D", "Gem. prijs E/F"]:
        display[col] = display[col].map(lambda x: "EUR {:,.0f}".format(x))
    display["Verschil mediaan"] = display["Verschil mediaan"].map(lambda x: "EUR {:+,.0f}".format(x))
    display["Verschil mediaan %"] = display["Verschil mediaan %"].map(lambda x: "{:+.1f}%".format(x))
    display["% E/F van totaal"] = display["% E/F van totaal"].map(lambda x: "{:.1f}%".format(x))
    display["n A-D"] = display["n A-D"].astype(int)
    display["n E/F"] = display["n E/F"].astype(int)
    display["n totaal"] = display["n totaal"].astype(int)
    print()
    print(display.to_string())
    print()


jaar_label = " ({})".format(args.jaar) if args.jaar else " (2023-2025)"
print_table(tbl_houses, "WONINGEN -- Prijs per EPC-groep per gemeente" + jaar_label)
print_table(tbl_apartments, "APPARTEMENTEN -- Prijs per EPC-groep per gemeente" + jaar_label)


# ============================================================
# Year-over-year comparison per municipality (2023 > 2024 > 2025)
# ============================================================
if not args.jaar:
    years = [2023, 2024, 2025]
    # Year-over-year uses A/B vs E/F to match article price premium framing
    yoy_good = ["A", "A+", "A++", "B"]
    yoy_poor = ["E", "F"]

    def build_yoy_table(data, property_type, min_n=5):
        d = data[data["property_type"] == property_type].copy()
        municipalities = sorted(d["municipality_name"].unique())

        rows = []
        for mun in municipalities + ["KUSTBREED", "KUST EXCL. KNOKKE"]:
            if mun == "KUSTBREED":
                dm = d
            elif mun == "KUST EXCL. KNOKKE":
                dm = d[d["municipality_name"] != "Knokke-Heist"]
            else:
                dm = d[d["municipality_name"] == mun]
            row = {}
            for y in years:
                dy = dm[dm["jaar"] == y]
                g = dy[dy["epc_label_display"].isin(yoy_good)]["price_value"]
                p = dy[dy["epc_label_display"].isin(yoy_poor)]["price_value"]
                row["A/B {}".format(y)] = g.median() if len(g) >= min_n else None
                row["n A/B {}".format(y)] = len(g)
                row["E/F {}".format(y)] = p.median() if len(p) >= min_n else None
                row["n E/F {}".format(y)] = len(p)
                if row["A/B {}".format(y)] and row["E/F {}".format(y)]:
                    row["Kloof {}".format(y)] = row["A/B {}".format(y)] - row["E/F {}".format(y)]
                else:
                    row["Kloof {}".format(y)] = None
            # Evolutie A/B
            for y in [2024, 2025]:
                prev = row.get("A/B {}".format(y - 1))
                curr = row.get("A/B {}".format(y))
                if prev and curr and prev > 0:
                    row["Evol A/B {}".format(y)] = ((curr - prev) / prev) * 100
                else:
                    row["Evol A/B {}".format(y)] = None
            # Evolutie E/F
            for y in [2024, 2025]:
                prev = row.get("E/F {}".format(y - 1))
                curr = row.get("E/F {}".format(y))
                if prev and curr and prev > 0:
                    row["Evol E/F {}".format(y)] = ((curr - prev) / prev) * 100
                else:
                    row["Evol E/F {}".format(y)] = None
            rows.append(row)

        tbl = pd.DataFrame(rows, index=municipalities + ["KUSTBREED", "KUST EXCL. KNOKKE"])
        tbl.index.name = "Gemeente"
        return tbl

    tbl_yoy_apt = build_yoy_table(df, "Apartment")
    tbl_yoy_house = build_yoy_table(df, "House")

    def print_yoy_table(tbl, label):
        print()
        print("=" * 180)
        print(label)
        print("=" * 180)
        display = tbl.copy()
        for col in display.columns:
            if col.startswith("n "):
                display[col] = display[col].astype(int)
            elif col.startswith("Evol"):
                display[col] = display[col].map(lambda x: "{:+.1f}%".format(x) if pd.notna(x) else "--")
            else:
                display[col] = display[col].map(lambda x: "EUR {:,.0f}".format(x) if pd.notna(x) else "--")
        # Select display columns
        cols = []
        for y in years:
            cols.extend(["A/B {}".format(y), "n A/B {}".format(y)])
            if y > 2023:
                cols.append("Evol A/B {}".format(y))
        for y in years:
            cols.extend(["E/F {}".format(y), "n E/F {}".format(y)])
            if y > 2023:
                cols.append("Evol E/F {}".format(y))
        for y in years:
            cols.append("Kloof {}".format(y))
        display = display[[c for c in cols if c in display.columns]]
        print()
        print(display.to_string())
        print()

    print_yoy_table(tbl_yoy_apt, "JAAR-OVER-JAAR APPARTEMENTEN -- Mediaanprijs A/B vs E/F per gemeente (2023 > 2024 > 2025)")
    print_yoy_table(tbl_yoy_house, "JAAR-OVER-JAAR WONINGEN -- Mediaanprijs A/B vs E/F per gemeente (2023 > 2024 > 2025)")


# ============================================================
# Excel export with blue table formatting
# ============================================================
if args.excel:
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    excel_path = os.path.join(RESULTS_DIR, args.excel)
    table_counter = 0

    def style_sheet(ws, tbl_df, has_index=True):
        """Apply blue table style and number formatting to a worksheet."""
        global table_counter
        table_counter += 1

        # Determine data range
        min_row, min_col = 1, 1
        max_row = ws.max_row
        max_col = ws.max_column
        ref = "A1:{}{}".format(get_column_letter(max_col), max_row)

        # Apply blue table style
        table = Table(displayName="Tabel{}".format(table_counter), ref=ref)
        style = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False
        )
        table.tableStyleInfo = style
        ws.add_table(table)

        # Get column names from header row
        headers = [ws.cell(row=1, column=c).value for c in range(1, max_col + 1)]

        # Apply number formats per column
        for col_idx, header in enumerate(headers, start=1):
            if header is None:
                continue
            h = str(header)
            # Determine format based on column name
            if any(k in h for k in ["Evol", "Verschil mediaan %", "Prijsverschil %", "% E/F"]):
                fmt = '0.0"%"'
            elif h.startswith("n ") or h == "n totaal" or h.startswith("n_"):
                fmt = '#,##0'
            else:
                # Price columns
                fmt = '#,##0'

            for row_idx in range(2, max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    cell.number_format = fmt

        # Auto-fit column widths (approximate)
        for col_idx in range(1, max_col + 1):
            max_len = len(str(headers[col_idx - 1] or ""))
            for row_idx in range(2, max_row + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len("{:,.0f}".format(val)) if isinstance(val, (int, float)) else len(str(val)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 20)

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        tbl_houses.to_excel(writer, sheet_name="Woningen", index_label="Gemeente")
        tbl_apartments.to_excel(writer, sheet_name="Appartementen", index_label="Gemeente")
        if not args.jaar:
            tbl_yoy_apt.to_excel(writer, sheet_name="YoY Appartementen")
            tbl_yoy_house.to_excel(writer, sheet_name="YoY Woningen")

        # Apply styling to all sheets
        for sheet_name in writer.sheets:
            style_sheet(writer.sheets[sheet_name], None)

    print("Exported to: {}".format(excel_path))
