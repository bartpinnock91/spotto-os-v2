import argparse
import os
import pandas as pd
import numpy as np
import warnings
warnings.filterwarnings("ignore")
pd.set_option("display.max_columns", 30)
pd.set_option("display.width", 200)
pd.set_option("display.max_rows", 100)

parser = argparse.ArgumentParser(description="Kustinventaris EPC analyse")
parser.add_argument("--excel", type=str, nargs="?", const="analyse_kustinventaris.xlsx", default=None,
                    help="Export to Excel. Optionally specify filename (default: results/analyse_kustinventaris.xlsx)")
args = parser.parse_args()

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(SCRIPT_DIR, "source", "data_kustinventaris.csv")
RESULTS_DIR = os.path.join(SCRIPT_DIR, "results")
df = pd.read_csv(CSV, na_values=["null", "NULL", "Null", ""], encoding="latin-1")
print("Total rows loaded: {:,}".format(len(df)))

df_res = df[df["is_residential"] == True].copy()
print("After is_residential=TRUE: {:,}".format(len(df_res)))
df_sale = df_res[df_res["transaction_type"] == "Sale"].copy()
print("After transaction_type=Sale: {:,}".format(len(df_sale)))

df_sale["price_value"] = pd.to_numeric(df_sale["price_value"], errors="coerce")
df_sale["construction_square_meters"] = pd.to_numeric(df_sale["construction_square_meters"], errors="coerce")

def apply_quality_filters(data, pt):
    d = data[data["property_type"] == pt].copy()
    print()
    print("--- Filtering for {} ---".format(pt))
    print("  property_type={}: {:,}".format(pt, len(d)))
    d = d[d["price_value"] > 50000]
    print("  price > 50000: {:,}".format(len(d)))
    d = d[d["construction_square_meters"] > 20]
    print("  sqm > 20: {:,}".format(len(d)))
    d = d[d["epc_label_display"].notna()]
    d = d[d["epc_label_display"] != "Unknown"]
    print("  epc valid: {:,}".format(len(d)))
    d["price_per_sqm"] = d["price_value"] / d["construction_square_meters"]
    d = d[(d["price_per_sqm"] >= 500) & (d["price_per_sqm"] <= 15000)]
    print("  price/sqm 500-15000: {:,}".format(len(d)))
    return d

df_apt = apply_quality_filters(df_sale, "Apartment")
df_house = apply_quality_filters(df_sale, "House")

good_epc = ["A++", "A+", "A", "B"]
poor_epc = ["E", "F"]
all_labels = ["A", "B", "C", "D", "E", "F"]

def fmt_eur(x): return "EUR {:,.0f}".format(x)
def fmt_eur_s(x): return "EUR {:+,.0f}".format(x)
def fmt_pct(x): return "{:+.1f}%".format(x)

def good_vs_poor(data, min_poor=10):
    good = data[data["epc_label_display"].isin(good_epc)].groupby("municipality_name")["price_per_sqm"].agg(["median", "count"])
    good.columns = ["median_good", "n_good"]
    poor = data[data["epc_label_display"].isin(poor_epc)].groupby("municipality_name")["price_per_sqm"].agg(["median", "count"])
    poor.columns = ["median_poor", "n_poor"]
    merged = good.join(poor, how="inner")
    merged = merged[merged["n_poor"] >= min_poor]
    merged["diff_abs"] = merged["median_good"] - merged["median_poor"]
    merged["diff_pct"] = (merged["diff_abs"] / merged["median_poor"]) * 100
    merged = merged.sort_values("diff_pct", ascending=False)
    merged["n_good"] = merged["n_good"].astype(int)
    merged["n_poor"] = merged["n_poor"].astype(int)
    return merged

def fmt_good_vs_poor(merged):
    r = merged.copy()
    r["median_good"] = r["median_good"].map(fmt_eur)
    r["median_poor"] = r["median_poor"].map(fmt_eur)
    r["diff_abs"] = r["diff_abs"].map(fmt_eur_s)
    r["diff_pct"] = r["diff_pct"].map(fmt_pct)
    return r

def coastwide(data, prop_label):
    ext = ["A++", "A+"] + all_labels
    results = []
    for epc in ext:
        sub = data[data["epc_label_display"] == epc]
        if len(sub) > 0:
            results.append({
                "EPC Label": epc,
                "Median EUR/m2": round(sub["price_per_sqm"].median(), 0),
                "Mean EUR/m2": round(sub["price_per_sqm"].mean(), 0),
                "n": len(sub)
            })
    rdf = pd.DataFrame(results)

    # Good vs Poor summary row
    gd = data[data["epc_label_display"].isin(good_epc)]
    pr = data[data["epc_label_display"].isin(poor_epc)]
    if len(gd) > 0 and len(pr) > 0:
        gm = gd["price_per_sqm"].median()
        pm = pr["price_per_sqm"].median()
        diff = gm - pm
        pct = (diff / pm) * 100
        summary = pd.DataFrame([
            {"EPC Label": "Good (A/A+/A++/B)", "Median EUR/m2": round(gm, 0), "Mean EUR/m2": None, "n": len(gd)},
            {"EPC Label": "Poor (E/F)", "Median EUR/m2": round(pm, 0), "Mean EUR/m2": None, "n": len(pr)},
            {"EPC Label": "Difference", "Median EUR/m2": round(diff, 0), "Mean EUR/m2": None, "n": "{:+.1f}%".format(pct)},
        ])
        rdf = pd.concat([rdf, pd.DataFrame([{"EPC Label": "---"}]), summary], ignore_index=True)

    return rdf

# ============================================================
# A) APARTMENTS - Good vs Poor by municipality
# ============================================================
print()
print("=" * 110)
print("A) APARTMENTS - Median price/m2 by EPC quality per municipality")
print("   Good EPC = A, A+, A++, B  |  Poor EPC = E, F")
print("=" * 110)

result_a = good_vs_poor(df_apt, min_poor=10)
print()
print("Municipalities with n_poor >= 10:")
print()
print(fmt_good_vs_poor(result_a).to_string())

# ============================================================
# B) APARTMENTS - By individual EPC label per municipality
# ============================================================
print()
print()
print("=" * 110)
print("B) APARTMENTS - Median price/m2 by individual EPC label per municipality")
print("   Only municipalities with total n >= 50")
print("=" * 110)

apt_labels = df_apt[df_apt["epc_label_display"].isin(all_labels)]
pivot_med = apt_labels.pivot_table(values="price_per_sqm", index="municipality_name", columns="epc_label_display", aggfunc="median").reindex(columns=all_labels)
pivot_cnt = apt_labels.pivot_table(values="price_per_sqm", index="municipality_name", columns="epc_label_display", aggfunc="count").reindex(columns=all_labels)
total_n = pivot_cnt.sum(axis=1)
valid_mun = total_n[total_n >= 50].index
pmed = pivot_med.loc[valid_mun]
pcnt = pivot_cnt.loc[valid_mun]

# Build a combined table with median and count
result_b_display = pd.DataFrame(index=valid_mun)
for label in all_labels:
    vals = []
    for mu in valid_mun:
        med = pmed.loc[mu, label]
        cnt = pcnt.loc[mu, label]
        if pd.notna(med) and pd.notna(cnt):
            vals.append("EUR {:,.0f} (n={})".format(med, int(cnt)))
        else:
            vals.append("--")
    result_b_display[label] = vals
result_b_display["Total_n"] = total_n.loc[valid_mun].astype(int)
result_b_display = result_b_display.sort_values("Total_n", ascending=False)

# Raw numeric table for Excel
result_b_raw = pmed.copy()
result_b_raw["Total_n"] = total_n.loc[valid_mun].astype(int)
result_b_raw = result_b_raw.sort_values("Total_n", ascending=False)

print()
print(result_b_display.to_string())

# ============================================================
# C) HOUSES - Good vs Poor by municipality
# ============================================================
print()
print()
print("=" * 110)
print("C) HOUSES - Median price/m2 by EPC quality per municipality")
print("   Good EPC = A, A+, A++, B  |  Poor EPC = E, F")
print("=" * 110)

result_c = good_vs_poor(df_house, min_poor=10)
if len(result_c) == 0:
    print("No municipalities with n_poor >= 10. Trying n_poor >= 5:")
    result_c = good_vs_poor(df_house, min_poor=5)
if len(result_c) == 0:
    print("Showing all (n_poor >= 1):")
    result_c = good_vs_poor(df_house, min_poor=1)
print()
print(fmt_good_vs_poor(result_c).to_string())

# ============================================================
# D) COAST-WIDE - All municipalities combined
# ============================================================
print()
print()
print("=" * 110)
print("D) COAST-WIDE - Median price/m2 per EPC label (all municipalities combined)")
print("=" * 110)

result_d_apt = coastwide(df_apt, "APARTMENTS")
print()
print("--- APARTMENTS ---")
print(result_d_apt.to_string(index=False))

result_d_house = coastwide(df_house, "HOUSES")
print()
print("--- HOUSES ---")
print(result_d_house.to_string(index=False))

print()
print("=" * 110)
print("Analysis complete.")
print("=" * 110)

# ============================================================
# Excel export
# ============================================================
if args.excel:
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    excel_path = os.path.join(RESULTS_DIR, args.excel)
    table_counter = 0

    def style_sheet(ws):
        """Apply blue table style and number formatting to a worksheet."""
        global table_counter
        table_counter += 1

        max_row = ws.max_row
        max_col = ws.max_column
        ref = "A1:{}{}".format(get_column_letter(max_col), max_row)

        table = Table(displayName="Tabel{}".format(table_counter), ref=ref)
        style = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False
        )
        table.tableStyleInfo = style
        ws.add_table(table)

        headers = [ws.cell(row=1, column=c).value for c in range(1, max_col + 1)]

        for col_idx, header in enumerate(headers, start=1):
            if header is None:
                continue
            h = str(header)
            if any(k in h for k in ["diff_pct", "pct"]):
                fmt = '0.0"%"'
            elif h.startswith("n") or h == "Total_n":
                fmt = '#,##0'
            else:
                fmt = '#,##0'

            for row_idx in range(2, max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    cell.number_format = fmt

        for col_idx in range(1, max_col + 1):
            max_len = len(str(headers[col_idx - 1] or ""))
            for row_idx in range(2, max_row + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len("{:,.0f}".format(val)) if isinstance(val, (int, float)) else len(str(val)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 20)

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        result_a.to_excel(writer, sheet_name="Apt goed vs zwak")
        result_b_raw.to_excel(writer, sheet_name="Apt per EPC label")
        result_c.to_excel(writer, sheet_name="Huizen goed vs zwak")
        result_d_apt.to_excel(writer, sheet_name="Kustbreed apt", index=False)
        result_d_house.to_excel(writer, sheet_name="Kustbreed huizen", index=False)

        for sheet_name in writer.sheets:
            style_sheet(writer.sheets[sheet_name])

    print()
    print("Exported to: {}".format(excel_path))
